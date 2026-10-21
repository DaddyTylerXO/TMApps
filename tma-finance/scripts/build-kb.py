"""Build the KB Word draft from the same article source used by the app.
Run with the Codex bundled Python and Node runtimes, then render and inspect it.
"""
import argparse,json,subprocess
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

def link(paragraph,label,target,internal=False):
    h=OxmlElement('w:hyperlink')
    if internal:h.set(qn('w:anchor'),target)
    else:h.set(qn('r:id'),paragraph.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'163A2B');pr.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);paragraph._p.append(h)

def bookmark(p,name,number):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number));p._p.insert(0,start);p._p.append(end)

args=argparse.ArgumentParser();args.add_argument('--node',required=True);args.add_argument('--output',required=True);a=args.parse_args()
root=Path(__file__).resolve().parents[1]
script='import {kbArticles} from "./dist/kb.js"; process.stdout.write(JSON.stringify(kbArticles));'
articles=json.loads(subprocess.check_output([a.node,'--input-type=module','-e',script],cwd=root).decode('utf-8'))
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name,size in [('Normal',10.5),('Title',27),('Heading 1',19),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
p=doc.add_paragraph('TMA Finance KB Library','Title');bookmark(p,'index',0)
doc.add_paragraph('Version 0.1.3 | Updated 21 October 2026')
doc.add_paragraph('Use this library to choose a market connection, set up a personal API key and understand the data available in TMA-Finance. Shared Alpha Vantage is the default selection and is awaiting backend deployment. Personal connections are available now, subject to account permissions.')
doc.add_heading('Guide index',1)
for article in articles:link(doc.add_paragraph(),article['id']+'  '+article['title'],article['id'].replace('-','_'),True)
doc.add_paragraph('Provider prices, quotas and permissions may change. Check the linked official pages before selecting a plan. The app does not purchase subscriptions or upgrade accounts. No real credentials are included in this library.')
for i,article in enumerate(articles,1):
    doc.add_page_break();p=doc.add_paragraph(article['id'].replace('-',' ')+' '+article['title'],'Heading 1');bookmark(p,article['id'].replace('-','_'),i)
    link(doc.add_paragraph(),'Back to KB index','index',True)
    for heading,text in article['sections']:
        doc.add_heading(heading,2);doc.add_paragraph(text)
    if article['id']=='Finance-KB1':
        p=doc.add_paragraph();r=p.add_run();r.add_picture(str(root/'docs/green-expense-dialog.png'),width=Inches(3.6))
        r._r.xpath('.//wp:docPr')[0].set('descr','Empty expense dialog in the forest green desktop interface')
        doc.add_paragraph('An empty expense dialog from an isolated test profile. No personal finance records are shown.')
    for label,url in article['links']:link(doc.add_paragraph(),label,url)
doc.core_properties.author='TMApps';doc.core_properties.last_modified_by='TMApps';doc.core_properties.title='TMA Finance KB Library'
output=Path(a.output);output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
print('Created KB Word draft with '+str(len(articles))+' articles; render review required.')
