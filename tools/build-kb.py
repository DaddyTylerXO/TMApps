"""Run with the bundled workspace Python. Render and inspect DOCX before release."""
import json
import re
from io import BytesIO
from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
content = json.loads((root / 'content/experience.json').read_text())
version = json.loads((root / 'desktop/package.json').read_text())['version']
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.85)
for name in ['Normal', 'Body Text']:
    style = document.styles[name]
    style.font.name, style.font.size = 'Calibri', Pt(11)
    style.paragraph_format.space_after = Pt(8)
document.styles['Normal'].paragraph_format.line_spacing = 1.12
title = document.styles['Title']
title.font.name, title.font.size, title.font.bold = 'Calibri', Pt(27), True
title.font.color.rgb = RGBColor(0, 0, 0)
document.styles['Subtitle'].font.color.rgb = RGBColor(0, 0, 0)
document.styles['Caption'].font.color.rgb = RGBColor.from_string('555555')
for style in document.styles:
    for border in list(style.element.findall('.//' + qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Heading 1', 'Heading 2']:
    style = document.styles[name]
    style.font.name, style.font.size = 'Calibri', Pt(16)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True

bookmark_id = 0
targets = {article['id']: article['id'].replace('-', '_') for article in content['articles']}


def bookmark(paragraph, name):
    global bookmark_id
    bookmark_id += 1
    start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)
    end.set(qn('w:id'), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def link(paragraph, label, target, button=False):
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), target)
    hyperlink.set(qn('w:history'), '1')
    run, properties = OxmlElement('w:r'), OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), 'FFFFFF' if button else '9B4300')
    properties.append(color)
    if button:
        fill = OxmlElement('w:shd')
        fill.set(qn('w:fill'), 'CC5500')
        properties.append(fill)
        properties.append(OxmlElement('w:b'))
    else:
        underline = OxmlElement('w:u')
        underline.set(qn('w:val'), 'single')
        properties.append(underline)
    run.append(properties)
    text = OxmlElement('w:t')
    text.set(qn('xml:space'), 'preserve')
    # Keep a KB button on one line, including when it follows a long sentence.
    text.text = '  ' + label.replace('-', '\u2011') + '  ' if button else label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def referenced_text(paragraph, text):
    for part in re.split(r'(Launcher-KB\d+)', text):
        if part in targets:
            link(paragraph, part, targets[part], button=True)
        else:
            paragraph.add_run(part)


# Reuse real UI captures. Crop only the area explained by the article so text
# stays readable on a printed page; never substitute mockups for application UI.
illustrations = {
    'Launcher-KB1': ('launcher-preview.png', None, 6.6,
                     'The application library after opening the launcher. Available apps depend on your operating system.'),
    'Launcher-KB2': ('install-preview.png', None, 6.0,
                     'App installation offers a destination folder, startup and desktop shortcut options on Windows.'),
    'Launcher-KB3': ('profile-preview.png', (0, 0, 600, 509), 5.5,
                     'The optional profile shares suggestions with Mail. Provider authentication is still completed in the app.'),
    'Launcher-KB5': ('settings-preview.png', (0, 0, 600, 443), 5.8,
                     'Appearance controls with the Grey palette selected. Save keeps your chosen layout and colours.'),
    'Launcher-KB6': ('tray-menu-preview.png', None, 2.3,
                     'Choose Quit from the tray menu to release launcher memory.'),
    'Launcher-KB8': ('help-preview.png', (0, 0, 600, 540), 5.4,
                     'Help provides the KB Library and actions for GitHub support forms.'),
    'Launcher-KB9': ('details-preview.png', None, 5.4,
                     'App details show version information and release notes. Media slots remain unpublished until assets are supplied.'),
    'Launcher-KB10': ('uninstall-preview.png', None, 5.8,
                      'The uninstall dialog lets you choose whether the selected app keeps its standard local data.'),
    'Launcher-KB11': ('settings-search-preview.png', None, 5.8,
                      'Search by keywords or a KB number to find relevant knowledge articles.'),
}


def illustrate(article_id):
    if article_id not in illustrations:
        return
    filename, crop, width, caption = illustrations[article_id]
    with Image.open(root / 'docs' / filename) as original:
        picture = original.crop(crop) if crop else original.copy()
        stream = BytesIO()
        picture.save(stream, format='PNG')
        stream.seek(0)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(stream, width=Inches(width))
    paragraph = document.add_paragraph(caption, 'Caption')
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.space_after = Pt(12)
document.add_paragraph('TMApps Launcher KB Library', 'Title')
document.add_paragraph('User guides setup and troubleshooting', 'Subtitle')
document.add_paragraph(f'Release {version}  |  21 October 2026  |  DaddyTylerXO')
document.add_paragraph('Use this library to install the launcher, manage your applications, find settings, and report problems. Each article has a stable KB number that you can type into launcher Help or settings search. Shared TMApps sign-in is prepared for deployment and requires backend configuration before activation. App support varies by operating system.')
document.add_paragraph('Use the orange KB buttons to jump to an article, then Back to index to return. In Microsoft Word, hold Ctrl while clicking a link if required by your settings. Screenshots illustrate the launcher controls; versions and available apps may change.')
bookmark(document.add_heading('Article index', 1), 'KB_Index')
table = document.add_table(rows=1, cols=2)
table.style = 'Table Grid'
table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'KB number', 'Subject'
for article in content['articles']:
    cells = table.add_row().cells
    link(cells[0].paragraphs[0], article['id'], targets[article['id']], button=True)
    link(cells[1].paragraphs[0], article['title'], targets[article['id']])
table.autofit = False
table.columns[0].width, table.columns[1].width = Inches(1.65), Inches(5.15)
for i, row in enumerate(table.rows):
    for column, cell in enumerate(row.cells):
        cell.width = Inches(1.65 if column == 0 else 5.15)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'EDE8E2' if i == 0 else ('F5F5F5' if i % 2 else 'FFFFFF'))
        cell._tc.get_or_add_tcPr().append(shading)
        borders = OxmlElement('w:tcBorders')
        for edge in ['top', 'left', 'bottom', 'right']:
            border = OxmlElement('w:' + edge)
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:color'), 'D9D9D9')
            borders.append(border)
        cell._tc.get_or_add_tcPr().append(borders)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10)
                run.bold = i == 0
for article in content['articles']:
    # Brief related articles share their predecessor's page.
    if article['id'] not in ['Launcher-KB4', 'Launcher-KB7', 'Launcher-KB12']:
        document.add_page_break()
    heading = document.add_heading(article['id'].replace('-', ' ') + ' ' + article['title'], 1)
    bookmark(heading, targets[article['id']])
    reference = document.add_paragraph()
    link(reference, 'Back to index', 'KB_Index', button=True)
    reference.paragraph_format.keep_with_next = True
    sentences = re.split(r'(?<=[.!?])\s+(?=[A-Z])', article['body'])
    parts = [' '.join(sentences[i:i + 3]) for i in range(0, len(sentences), 3)]
    for i, part in enumerate(parts):
        paragraph = document.add_paragraph()
        referenced_text(paragraph, part)
        paragraph.paragraph_format.keep_with_next = i < len(parts) - 1
    illustrate(article['id'])
document.add_heading('Maintaining this library', 1)
document.add_paragraph('The in-app articles and this Word document are generated from content/experience.json. Keep each KB number stable when updating its instructions. Regenerate and visually inspect the Word document when article text changes. Deployment and app-integration details are in backend/README.md and docs/shared-profile.md on the launcher branch.')
footer = section.footer.paragraphs[0]
footer.text = 'TMApps Launcher KB Library  •  '
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
document.core_properties.title = 'TMApps Launcher KB Library'
document.core_properties.author = 'DaddyTylerXO'
document.save(root / 'docs/KB Library.docx')
