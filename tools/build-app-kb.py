"""Build a numbered Word KB from docs/kb.json using the bundled Python runtime."""
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

root = Path(__file__).resolve().parents[1]
config = json.loads((root / 'docs/kb-config.json').read_text(encoding='utf-8'))
articles = json.loads((root / 'docs/kb.json').read_text(encoding='utf-8'))
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.75)
s.left_margin = s.right_margin = Inches(.8)
for name, size in [('Normal', 11), ('Title', 26), ('Heading 1', 17), ('Heading 2', 13)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(8)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)

def bookmark(p, name, number):
    start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    start.set(qn('w:id'), str(number)); start.set(qn('w:name'), name)
    end.set(qn('w:id'), str(number)); p._p.insert(0, start); p._p.append(end)

def link(p, text, target, external=False):
    node, run, props = OxmlElement('w:hyperlink'), OxmlElement('w:r'), OxmlElement('w:rPr')
    if external:
        node.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    else:
        node.set(qn('w:anchor'), target)
    colour, shade, bold = OxmlElement('w:color'), OxmlElement('w:shd'), OxmlElement('w:b')
    colour.set(qn('w:val'), 'FFFFFF'); shade.set(qn('w:fill'), 'CC5500')
    props.extend([colour, shade, bold]); run.append(props)
    text_node = OxmlElement('w:t'); text_node.set(qn('xml:space'), 'preserve')
    text_node.text = '  ' + text.replace('-', '\u2011') + '  '
    run.append(text_node); node.append(run); p._p.append(node)

targets = {a['id']: a['id'].replace('-', '_') for a in articles}
header = s.header.paragraphs[0]
logo = root / config.get('logo', 'missing.png')
if logo.is_file(): header.add_run().add_picture(str(logo), width=Inches(.25))
header.add_run('  ' + config['name']).bold = True
footer = s.footer.paragraphs[0]
footer.add_run(config['name'] + ' KB Library  |  Source update  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.add_paragraph(config['name'].replace('-', ' ') + ' KB Library', 'Title')
doc.add_paragraph('User guides and setup information')
doc.add_paragraph('Source update  |  21 October 2026')
doc.add_paragraph('This library explains the available app features and the Spotify source integration prepared for its next release. Existing public downloads do not gain these new Spotify controls until a companion app update is published. No live Spotify account was used for verification.')
link(doc.add_paragraph(), 'TMApps master KB Library', 'https://github.com/DaddyTylerXO/TMApps/blob/TMApps-Launcher/docs/KB%20Library.docx', True)
bookmark(doc.add_paragraph('Article index', 'Heading 1'), 'KB_Index', 0)
for a in articles:
    p = doc.add_paragraph(); link(p, a['id'], targets[a['id']]); p.add_run('  ' + a['title'])
for i, a in enumerate(articles, 1):
    doc.add_page_break()
    bookmark(doc.add_paragraph(a['id'].replace('-', ' ') + ' ' + a['title'], 'Heading 1'), targets[a['id']], i)
    link(doc.add_paragraph(), 'Back to index', 'KB_Index')
    for text in a['body'].split('\n\n'):
        p = doc.add_paragraph()
        for part in re.split(r'([A-Za-z]+-KB\d+)', text):
            if part in targets: link(p, part, targets[part])
            else: p.add_run(part)
    image = root / a.get('image', 'missing.png')
    if image.is_file():
        with Image.open(image) as bitmap: width, height = bitmap.size
        scale = min(6.1 / width, 4.3 / height)
        picture = doc.add_picture(str(image), width=Inches(width * scale))
        picture._inline.docPr.set('descr', a.get('caption', a['title']))
        p = doc.add_paragraph(a.get('caption', 'Actual application view using a disposable local profile.'))
        for run in p.runs: run.font.size = Pt(9)
anchors = {n.get(qn('w:name')) for n in doc.element.xpath('//w:bookmarkStart')}
assert all(n.get(qn('w:anchor')) in anchors for n in doc.element.xpath('//w:hyperlink[@w:anchor]'))
doc.core_properties.title = config['name'].replace('-', ' ') + ' KB Library'
doc.core_properties.author = 'TMApps'
doc.save(root / config.get('output', 'docs/KB Library.docx'))
print(config['name'] + ': built ' + str(len(articles)) + ' linked KB articles')
