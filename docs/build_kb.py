"""Rebuild the branded KB using python-docx from the managed document runtime."""
import json
import re
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
GREEN = "285B3D"
registry = json.loads((DOCS / "kb-registry.json").read_text(encoding="utf-8"))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.8), Inches(.75)
section.left_margin = section.right_margin = Inches(.8)
section.header_distance = section.footer_distance = Inches(.3)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.12
for name in ("Heading 1", "Heading 2"):
    doc.styles[name].font.color.rgb = RGBColor.from_string(GREEN)
doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for name in ("Subtitle", "Caption"):
    doc.styles[name].font.color.rgb = RGBColor.from_string(GREEN)
for border in list(doc.styles.element.xpath(".//w:pBdr")):
    border.getparent().remove(border)
header = section.header.paragraphs[0]
header.add_run().add_picture(str(ROOT / "assets" / "TMA-Mail.png"), width=Inches(.28))
run = header.add_run("  TMA-Mail  |  KB Library")
run.font.bold, run.font.color.rgb, run.font.size = True, RGBColor.from_string(GREEN), Pt(10)
footer = section.footer.paragraphs[0]
footer.add_run("TMA-Mail  •  23 October 2026  •  Page ").font.size = Pt(9)
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)

def bookmark(paragraph, name, number):
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(number)); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start); paragraph._p.append(end)

def link(paragraph, label, destination, internal=False):
    element = OxmlElement("w:hyperlink")
    if internal: element.set(qn("w:anchor"), destination)
    else: element.set(qn("r:id"), paragraph.part.relate_to(destination, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r"); properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), GREEN); properties.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); properties.append(underline)
    run.append(properties); text = OxmlElement("w:t"); text.text = label; run.append(text); element.append(run); paragraph._p.append(element)

def prose(text):
    paragraph = doc.add_paragraph()
    pattern = r"\[([^\]]+)\]\(([^)]+)\)|`([^`]+)`"
    position = 0
    for match in re.finditer(pattern, text):
        paragraph.add_run(text[position:match.start()])
        if match.group(1):
            target = match.group(2)
            article = next((a for a in registry["articles"] if a["file"] == target), None)
            link(paragraph, match.group(1), article["bookmark"] if article else "https://github.com/DaddyTylerXO/TMApps/blob/TMA-Mail/" + target.removeprefix("../") if not target.startswith("https://") else target, internal=article is not None)
        else:
            run = paragraph.add_run(match.group(3)); run.font.name = "Consolas"; run.font.size = Pt(9.5)
        position = match.end()
    paragraph.add_run(text[position:])
    return paragraph

doc.add_paragraph("TMA-Mail KB Library", "Title")
doc.add_paragraph("Cross platform migration preview", "Subtitle")
doc.add_paragraph("Updated 23 October 2026\nPublished Windows app 0.16.5\nMigration source preview 0.17.0-preview.1")
doc.add_paragraph("This library explains the sample-only desktop preview, how to build it, and its privacy and platform limits. Continue using the supported Windows app for real mail. macOS and Linux production compatibility is not complete.")
doc.add_heading("Guide index", level=1)
bookmark(doc.paragraphs[-1], "Mail_Index", 1)
for article in registry["articles"]:
    link(doc.add_paragraph(), article["id"] + "  " + article["title"], article["bookmark"], True)
doc.add_heading("Suite documentation", level=1)
doc.add_paragraph("The TMApps Master KB Library has not yet been published on the main branch at the reviewed revision. This library does not claim that an unavailable master document exists.")
link(doc.add_paragraph(), "Open TMApps main branch documentation", "https://github.com/DaddyTylerXO/TMApps/tree/main")

for number, article in enumerate(registry["articles"], 10):
    doc.add_page_break()
    heading = doc.add_heading(article["id"] + " " + article["title"], level=1)
    bookmark(heading, article["bookmark"], number)
    text = (DOCS / article["file"]).read_text(encoding="utf-8")
    chunks = text.split("\n\n")[1:]
    for chunk in chunks:
        chunk = chunk.strip()
        if not chunk or chunk.startswith("[Back to index]"): continue
        if chunk.startswith("!["):
            doc.add_picture(str(DOCS / "screenshots" / "migration-preview.png"), width=Inches(6.75))
            caption = doc.add_paragraph("Figure 1  Actual preview interface using synthetic mail", "Caption")
            caption.paragraph_format.space_after = Pt(10)
        elif chunk.startswith("## "): doc.add_heading(chunk[3:], level=2)
        else: prose(chunk.replace("\n", " "))
    link(doc.add_paragraph(), "Back to index", "Mail_Index", True)

doc.core_properties.title = "TMA-Mail KB Library"
doc.core_properties.subject = "Cross-platform migration preview guides"
doc.core_properties.author = "TMA-Mail"
doc.core_properties.last_modified_by = "TMA-Mail"
output = DOCS / "TMA-Mail KB Library.docx"
doc.save(output)
shutil.copyfile(output, DOCS / "KB Library.docx")

# Validate every internal navigation destination before rendering.
from zipfile import ZipFile
from lxml import etree
with ZipFile(output) as package:
    tree = etree.fromstring(package.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    targets = set(tree.xpath("//w:bookmarkStart/@w:name", namespaces=ns))
    anchors = tree.xpath("//w:hyperlink/@w:anchor", namespaces=ns)
    assert anchors and all(a in targets for a in anchors), "Broken KB navigation target"
print("Created canonical and compatibility KB libraries; internal links validated.")
