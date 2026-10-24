from pathlib import Path
import json, shutil, argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
root=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser();parser.add_argument('--screenshots',type=Path);args=parser.parse_args()
version=json.loads((root/'package.json').read_text(encoding='utf-8'))['version']
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)

for element in list(doc.styles.element.iter(qn('w:pBdr'))): element.getparent().remove(element)
doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
doc.styles['Subtitle'].font.color.rgb=RGBColor.from_string('595959')
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('595959')
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(11);style.paragraph_format.space_after=Pt(9)
for n in ['Heading 1','Heading 2']:
 doc.styles[n].font.color.rgb=RGBColor.from_string('7048A4')
header=sec.header.paragraphs[0];header.add_run().add_picture(str(root/'assets/icon.png'),width=Inches(.28));r=header.add_run('  TMA-Safe    KB Library');r.bold=True;r.font.color.rgb=RGBColor.from_string('7048A4')
footer=sec.footer.paragraphs[0];r=footer.add_run(f'TMA-Safe {version}  |  24 October 2026  |  ');r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('7048A4')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def bookmark(p,name,num):
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(num));start.set(qn('w:name'),name);end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(num));p._p.insert(0,start);p._p.append(end)
def link(p,text,target,external=False):
 h=OxmlElement('w:hyperlink')
 if external:h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
 else:h.set(qn('w:anchor'),target)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'7048A4');pr.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
articles=json.loads((root/'docs/kb.json').read_text(encoding='utf-8'))
doc.add_paragraph('TMA Safe KB Library','Title');doc.add_paragraph('Setup and practical help for your local password vault','Subtitle')
doc.add_paragraph('Use these guides to create a vault, organise passwords, adjust locking and recover safely from common errors. The app stores an encrypted file on your device; it has no cloud account or password reset.')
p=doc.add_paragraph();link(p,'TMA-Safe project branch','https://github.com/DaddyTylerXO/TMApps/tree/TMA-Safe',True)
doc.add_paragraph('TMApps Master KB Library: publication has not been verified. This child library is ready for registration when the master index is available.');p=doc.add_paragraph();link(p,'TMApps main repository','https://github.com/DaddyTylerXO/TMApps/tree/main',True)
p=doc.add_heading('Guide index',1);bookmark(p,'Index',0)
for a in articles:link(doc.add_paragraph(),a['id']+'  '+a['title'],a['id'].replace('-',''))
doc.add_heading('Before you begin',1)
doc.add_paragraph('Use a unique master passphrase and keep an encrypted backup. Store a separate key file away from the vault. Real vaults, key files, browser exports and pairing codes must never be attached to support reports.')
doc.add_paragraph('Core vault operations work offline. Browser filling also requires a locally installed companion extension and approval in the unlocked app. GitHub help pages and application downloads require connectivity. No online breach monitoring is enabled.')
for i,a in enumerate(articles,1):
 doc.add_page_break();p=doc.add_heading(a['id']+'  '+a['title'],1);bookmark(p,a['id'].replace('-',''),i);link(doc.add_paragraph(),'Back to index','Index')
 if a['image']:
  folder=root/'docs/images';folder.mkdir(exist_ok=True);im=Image.open((args.screenshots if args.screenshots else folder)/a['image'])
  if args.screenshots and a['image']=='create.png':im=im.crop((620,45,1190,760))
  elif args.screenshots and a['image']=='api-editor.png':im=im.crop((330,35,903,780))
  dest=folder/a['image'];im.save(dest)
  width=Inches(3.15 if a['image'] in ['create.png','api-editor.png'] else 6.45);doc.add_picture(str(dest),width=width)
  p=doc.add_paragraph(a['caption']);p.style='Caption'
 for text in a['paragraphs']:
  p=doc.add_paragraph()
  import re
  start=0
  for m in re.finditer(r'Safe-KB[1-9][0-9]*',text):
   p.add_run(text[start:m.start()]);link(p,m.group(),m.group().replace('-',''));start=m.end()
  p.add_run(text[start:])
doc.core_properties.author='TMApps';doc.core_properties.title='TMA Safe KB Library';doc.core_properties.subject='Offline password vault user guides';doc.core_properties.comments=''
out=root/'docs/TMA-Safe KB Library.docx';doc.save(out);shutil.copyfile(out,root/'docs/KB Library.docx');print(out)
