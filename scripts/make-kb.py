"""Create the Word guide from the same numbered articles as the offline UI.

Run with the Codex bundled Python after exporting .integration/kb.json and
capturing synthetic app screenshots. Render and visually inspect after creation.
"""
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from shutil import copyfile

root=Path(__file__).resolve().parents[1]
articles=json.loads((root/'.integration/kb.json').read_text(encoding='utf-8'))
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.27);section.page_height=Inches(11.69)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Header','Footer']:
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(9)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(30)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(14)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
for node in list(doc.styles.element.xpath('.//w:pBdr')):node.getparent().remove(node)
doc.core_properties.title='KB Library'
doc.core_properties.subject='TMA-Padlock user guides and troubleshooting'
doc.core_properties.author='TMApps'
doc.core_properties.keywords='Padlock-KB1, authenticator, offline, backup, security'
header=section.header.paragraphs[0]
header.add_run().add_picture(str(root/'assets/icon.png'),width=Inches(.28))
header.add_run('  TMA-Padlock  |  KB Library').bold=True
header.paragraph_format.space_after=Pt(8)
for r in header.runs:r.font.color.rgb=RGBColor(0,0,0);r.font.size=Pt(10)
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('TMA-Padlock 0.4.0  |  ').font.color.rgb=RGBColor(0,0,0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
bookmark_id=0
def bookmark(p,name):
    global bookmark_id
    bookmark_id+=1
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_id));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_id))
    p._p.insert(0,start);p._p.append(end)
def link(p,text,anchor=None,url=None):
    h=OxmlElement('w:hyperlink')
    if anchor:h.set(qn('w:anchor'),anchor)
    if url:
        from docx.opc.constants import RELATIONSHIP_TYPE
        h.set(qn('r:id'),p.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'832B50');pr.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');pr.append(underline)
    r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def referenced_paragraph(text):
    p=doc.add_paragraph()
    for segment in re.split(r'(Padlock-KB\d+)',text):
        if re.fullmatch(r'Padlock-KB\d+',segment):link(p,segment,anchor=segment.replace('-',''))
        else:p.add_run(segment)
    return p

def picture(name,width=6.1):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(root/'docs/screenshots'/name),width=Inches(width))
    p.paragraph_format.space_after=Pt(4)
def caption(text):
    p=doc.add_paragraph(text);p.paragraph_format.space_after=Pt(10)
    for r in p.runs:r.italic=True;r.font.size=Pt(9)
def article(i):
    a=articles[i-1];heading=doc.add_heading(a['id']+' '+a['title'],level=1);bookmark(heading,a['id'].replace('-',''))
    # Short paragraphs are easier to scan in the Word guide.
    text=a['text'].replace('’',"'").replace('–','-')
    sentences=text.split('. ');middle=max(1,len(sentences)//2)
    for chunk in [sentences[:middle],sentences[middle:]]:
        if chunk:referenced_paragraph('. '.join(chunk)+('' if chunk[-1].endswith('.') else '.'))
    p=doc.add_paragraph();link(p,'Back to index',anchor='PadlockIndex');p.paragraph_format.space_after=Pt(6)

doc.add_heading('KB Library',0)
doc.add_paragraph('TMA-Padlock 0.4.0 | User guides and setup information',style='Subtitle')
doc.add_paragraph('23 October 2026')
doc.add_paragraph('Suite reference: the TMApps Master KB Library was not yet published on main at verification. This guide preserves Padlock KB IDs for future registry links.')
doc.add_paragraph('Keep this guide with your recovery plan. All example accounts and codes in the screenshots are fictional. The same numbered articles are available offline in Settings and can be found with Settings search.')
bookmark(doc.add_heading('Find a guide',1),'PadlockIndex')
for a in articles:
    p=doc.add_paragraph();link(p,a['id']+'  '+a['title'],anchor=a['id'].replace('-',''));p.paragraph_format.space_after=Pt(4)
picture('tour.png',4.3)
caption('The separate demo introduces the app without changing real accounts.')
doc.add_page_break()
article(1);article(2)
picture('locked.png')
caption('First setup: choose vault protection and enter your master password.')
doc.add_page_break()
article(3);article(4)
picture('vault.png')
caption('Main authenticator list: expiry rings, code boxes and account menus.')
doc.add_page_break()
article(5);article(6)
table=doc.add_table(rows=1,cols=2)
for cell,name in zip(table.rows[0].cells,['mini.png','mini-menu.png']):
    cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraphs[0].add_run().add_picture(str(root/'docs/screenshots'/name),width=Inches(2.55))
caption('Compact window and right-click account menu, with synthetic accounts.')
doc.add_page_break()
article(7)
picture('appearance.png')
caption('Settings search finds appearance scheduling and transition duration.')
doc.add_paragraph('For an instant change, set Transition to 0. For a gentler scheduled change, try 2 or 30 minutes. To reduce memory use, close the mini window when it is not needed. The main app remains available through the tray.')
doc.add_page_break()
article(8);article(9)
doc.add_heading('Before sending a report',1)
for text in ['Review every word in the description. Never include setup keys, master passwords, one-time codes or encrypted vault files.','The optional diagnostic contains only app/runtime metadata, timestamps and generic failure categories. Raw exception messages and paths are omitted.','Upload opens a prefilled GitHub issue. Complete sign-in and select Submit new issue on GitHub; nothing is posted automatically.']:
    doc.add_paragraph(text,style='List Bullet')
doc.add_heading('Online reference',1)
doc.add_paragraph('github.com/DaddyTylerXO/TMApps/tree/TMA-Padlock/docs')
doc.add_paragraph('For the security assessment and performance measurements, see security-assessment.md and performance.md in that folder. Related guides: Padlock-KB3 for offline operation, Padlock-KB4 for recovery and Padlock-KB8 for security boundaries.')
doc.add_page_break()
article(10)
doc.add_paragraph('Keep the original encrypted vault and known-good backups safe while troubleshooting. Never upload a whole application-data folder. When contacting support, include the app version, operating system and architecture, the KB steps you tried and only sanitized error details.')
out=root/'docs/TMA-Padlock KB Library.docx'
doc.save(out)
copyfile(out,root/'docs/KB Library.docx')
print(out)
