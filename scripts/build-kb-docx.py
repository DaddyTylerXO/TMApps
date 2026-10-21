import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

entries = json.loads(Path('assets/kb.json').read_text(encoding='utf-8'))
doc = Document()
bookmark_id = 1
def anchor_for(value):
    return value.replace('-', '_')
def bookmark(paragraph, name):
    global bookmark_id
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), anchor_for(name))
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bookmark_id))
    bookmark_id += 1
    paragraph._p.insert(0, start)
    paragraph._p.append(end)
def link(paragraph, text, target, external=False):
    hyperlink = OxmlElement('w:hyperlink')
    if external:
        hyperlink.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    else:
        hyperlink.set(qn('w:anchor'), anchor_for(target))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    colour = OxmlElement('w:color')
    colour.set(qn('w:val'), '163D6A' if external else '332B12')
    properties.append(colour)
    if not external:
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'F4E8BA')
        properties.append(shade)
    run.append(properties)
    content = OxmlElement('w:t')
    content.set(qn('xml:space'), 'preserve')
    content.text = text
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
def body_text(paragraph, value):
    cursor = 0
    for match in re.finditer(r'https://[^\s]+', value):
        paragraph.add_run(value[cursor:match.start()])
        link(paragraph, match.group(), match.group(), external=True)
        cursor = match.end()
    paragraph.add_run(value[cursor:])
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.12
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.add_paragraph('TMA Socials Manager KB Library', 'Title')
doc.add_paragraph('User guides and setup information for version 0.4.0')
doc.add_paragraph('Use this library to configure connections, work with dashboards and drafts, adjust performance and appearance, and report issues safely. Each Socials KB number matches an article in the app. Credentials belong only in encrypted local settings; never include them in a public report.')
index = doc.add_paragraph('Article index', 'Heading 1')
bookmark(index, 'KB_Index')
for entry in entries:
    link(doc.add_paragraph(), '  ' + entry['id'] + '  ' + entry['title'] + '  ', entry['id'])
figures = {
    'Socials-KB5': ('widget-gallery.png', 'Visual display choices in Add widget. The dropdown remains available below this gallery.'),
    'Socials-KB6': ('image-draft.png', 'A sanitized example image attached to a local draft. Upload and Remove controls are above the preview.'),
    'Socials-KB7': ('discovery.png', 'YouTube discovery before its first requested check. The app distinguishes provider data from official discovery links.'),
    'Socials-KB8': ('appearance.png', 'Navigation position is selected in Appearance and workspace. Left, Right, Top and Bottom are available.'),
    'Socials-KB12': ('kb-search.png', 'Settings search returns the matching numbered article when KB search is enabled.'),
}
for entry in entries:
    doc.add_page_break()
    heading = doc.add_paragraph(entry['title'], 'Heading 1')
    bookmark(heading, entry['id'])
    doc.add_paragraph(entry['id'])
    for paragraph in entry['body'].split('\n\n'):
        body_text(doc.add_paragraph(), paragraph)
    if entry['id'] in figures:
        filename, caption = figures[entry['id']]
        image = Path('docs/images') / filename
        with Image.open(image) as bitmap:
            width, height = bitmap.size
        scale = min(6.2 / width, 2.65 / height, 1 / 85)
        picture = doc.add_picture(str(image), width=Inches(width * scale))
        picture._inline.docPr.set('descr', caption)
        p = doc.add_paragraph(caption)
        p.paragraph_format.space_after = Pt(10)
        for run in p.runs:
            run.font.size = Pt(10)
    p = doc.add_paragraph()
    p.add_run('Related articles  ').bold = True
    for target in entry['related']:
        link(p, '  ' + target + '  ', target)
        p.add_run('  ')
    link(doc.add_paragraph(), '  Back to index  ', 'KB_Index')
doc.core_properties.title = 'TMA Socials Manager KB Library'
doc.core_properties.subject = 'User guides and setup information'
doc.core_properties.author = 'TMApps'
doc.core_properties.comments = ''
doc.save('docs/KB Library.docx')
print('Created Word KB Library with 12 numbered articles.')
