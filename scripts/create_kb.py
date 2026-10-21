"""Generate the local Word KB draft from the in-app offline articles.
The result must be rendered/inspected and master-linked before publication.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

root = Path(__file__).resolve().parents[1]
qa = root / '.qa' / 'kb'
qa.mkdir(parents=True, exist_ok=True)
logo = Image.new('RGB', (96, 96), '#97572B')
draw = ImageDraw.Draw(logo)
draw.rectangle((18, 27, 45, 42), fill='white')
draw.rectangle((18, 36, 78, 75), fill='white')
logo.save(qa / 'logo.png')
articles = re.findall(r'id: "(Libraries-KB\d+)",\s+title: "([^"]+)",\s+text: "([^"]+)"', (root / 'app/Help.tsx').read_text(encoding='utf-8'))
assert len(articles) == 7

d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
s.top_margin = s.bottom_margin = Inches(.75)
s.left_margin = s.right_margin = Inches(.8)
s.header_distance = s.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    d.styles[name].font.name = 'Calibri'
    d.styles[name].font.color.rgb = RGBColor(0, 0, 0)
d.styles['Normal'].font.size = Pt(11)
s.header.paragraphs[0].add_run().add_picture(str(qa / 'logo.png'), width=Inches(.27))
s.header.paragraphs[0].add_run('  TMA-Libraries').bold = True
s.footer.paragraphs[0].add_run('TMA-Libraries  |  0.2.0 development draft  |  Page ')
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); s.footer.paragraphs[0]._p.append(f)

def bookmark(p, name, number):
    a = OxmlElement('w:bookmarkStart'); a.set(qn('w:id'), str(number)); a.set(qn('w:name'), name)
    b = OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'), str(number))
    p._p.insert(0, a); p._p.append(b)

def link(text, anchor):
    p = d.add_paragraph(); h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), anchor)
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r); p._p.append(h)

p = d.add_paragraph('TMA Libraries KB Library', 'Title'); bookmark(p, 'Index', 0)
d.add_paragraph('Windows desktop development preview 0.2.0 dev 1 | 22 October 2026')
d.add_paragraph('Use this guide to create and preserve local documents, recover from errors and understand format limitations. This is a development draft; full editor and standards verification remains open.')
d.add_paragraph('Master KB reference pending confirmation. Project reference: https://github.com/DaddyTylerXO/TMApps. Do not publish this draft until the actual master KB destination and every rendered page have been checked.')
for number, (code, title, text) in enumerate(articles, 1):
    link(code + '  ' + title, code.replace('-', '_'))
for number, (code, title, text) in enumerate(articles, 1):
    d.add_page_break()
    p = d.add_heading(code.replace('-', ' ') + '  ' + title, 1)
    bookmark(p, code.replace('-', '_'), number)
    d.add_paragraph(text)
    if code == 'Libraries-KB1':
        d.add_paragraph('Run release/win-unpacked/TMA-Libraries.exe with its entire package directory present. The legacy build/TMA-Libraries.exe is the older interface. Only development dependencies require installation access to the internet.')
        image = root / '.qa/home.png'
        if image.exists():
            d.add_picture(str(image), width=Inches(6.5))
    if code == 'Libraries-KB2':
        d.add_paragraph('The new user-data folder is normally %APPDATA%/TMA-Libraries/library-v2. Preserve the whole folder including managed assets when recovering. Independent dated native archives are recommended. Never add these files to GitHub.')
    if code == 'Libraries-KB4':
        d.add_paragraph('All-editor memory and final CPU samples exceed the initial targets on the development machine. Larger files, complex sheets and short save intervals can increase resource use. Full screen-reader and high-DPI verification is pending.')
    if code == 'Libraries-KB6':
        d.add_paragraph('Spreadsheet chart preview is not full XLSX chart export. Combined Book section exports, advanced layout fidelity and complete interactive round-trip verification are pending. Compatibility warnings must be reviewed before using exported copies.')
    link('Back to index', 'Index')
d.core_properties.author = 'TMApps'
d.core_properties.title = 'TMA Libraries KB Library'
d.core_properties.subject = 'Offline setup recovery and format guidance 0.2.0 development draft'
d.save(root / 'docs/TMA-Libraries KB Library.docx')
print('Created local Word KB draft with seven stable article IDs and internal index links; rendering and master reference pending.')
